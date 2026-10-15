"""Lectores de los Excel. Cada uno devuelve {tabla: (columnas, filas)} listo para cargar a SQL,
más un diccionario de datos sueltos ("extra") cuando el Excel trae algún parámetro útil."""
from pathlib import Path

import openpyxl

from .excel import buscar_encabezado, fecha, fecha_hora, norm, numero, texto, ubicar_columnas


def _hoja(wb, *contiene):
    for ws in wb.worksheets:
        nombre = norm(ws.title)
        if all(norm(c) in nombre for c in contiene):
            return ws
    raise ValueError(f"No encontré la hoja '{' '.join(contiene)}' (hojas: {', '.join(wb.sheetnames)})")


def _filas(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _celda(fila, idx):
    return fila[idx] if idx is not None and idx < len(fila) else None


COLUMNAS_CONSOLIDADO = [
    "fila", "pedido_genesis", "pedido_sap", "orden", "zfer", "zhal", "acero", "vehiculo",
    "fecha_envio_planta", "tipo_pedido", "pieza", "corte", "material", "entrega_produccion",
    "puesto_trabajo", "prioridad", "plano", "puntos", "fecha_fabricacion", "operario", "estado", "mercado",
]

MAPA_CONSOLIDADO = {
    "pedido_genesis": ["Pedido genesis"], "pedido_sap": ["Pedido SAP"],
    "orden": ["Orden fabricacion", "Orden de fabricacion", "Orden"], "zfer": ["Zfer"], "zhal": ["Zhal"],
    "acero": ["Acero"], "vehiculo": ["Vehiculo"], "fecha_envio_planta": ["Fecha envio planta"],
    "tipo_pedido": ["Tipo pedido"], "pieza": ["Pieza"], "corte": ["Corte"], "material": ["Material"],
    "entrega_produccion": ["Entrega produccion", "Entrega a produccion"], "puesto_trabajo": ["Puesto trabajo", "Puesto de trabajo"],
    "prioridad": ["Prioridad"], "plano": ["Plano"], "puntos": ["Puntos"],
    "fecha_fabricacion": ["Fecha fabricacion", "Fecha de fabricacion"], "operario": ["Operario"],
    "estado": ["Estado"], "mercado": ["Mercado"],
}

COLUMNAS_RECHAZOS = ["fila", "orden", "zfer", "zhal", "acero", "vehiculo", "plano", "fecha"]
MAPA_RECHAZOS = {"orden": ["Orden fabricacion", "Orden"], "zfer": ["Zfer"], "zhal": ["Zhal"], "acero": ["Acero"],
                 "vehiculo": ["Vehiculo"], "plano": ["Plano"], "fecha": ["Fecha"]}


def leer_consolidado(ruta: Path):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = _filas(_hoja(wb, "CONSOLIDADO"))
        h, enc = buscar_encabezado(filas, ["Orden fabricacion", "Material", "Operario"])
        c = ubicar_columnas(enc, MAPA_CONSOLIDADO, obligatorias=[
            "orden", "acero", "vehiculo", "pieza", "corte", "material", "entrega_produccion",
            "puesto_trabajo", "fecha_fabricacion", "operario"])
        consolidado = []
        for i, r in enumerate(filas[h + 1:], start=h + 2):
            orden = texto(_celda(r, c["orden"]), 40)
            if orden is None:
                continue
            g = lambda campo, largo=None: texto(_celda(r, c.get(campo)), largo)
            operario = _celda(r, c["operario"])
            consolidado.append((
                i, g("pedido_genesis", 40), g("pedido_sap", 40), orden, g("zfer", 40), g("zhal", 40), g("acero", 80),
                g("vehiculo", 200), fecha_hora(_celda(r, c.get("fecha_envio_planta"))), g("tipo_pedido", 60),
                g("pieza", 40), g("corte", 80), g("material", 60), fecha(_celda(r, c["entrega_produccion"])),
                g("puesto_trabajo", 40), g("prioridad", 40), g("plano", 40), numero(_celda(r, c.get("puntos"))),
                fecha(_celda(r, c["fecha_fabricacion"])),
                None if operario is None or str(operario).strip() == "" else str(operario).strip()[:120],
                g("estado", 60), g("mercado", 60),
            ))

        filas_r = _filas(_hoja(wb, "RECHAZOS"))
        h, enc = buscar_encabezado(filas_r, ["Fecha", "Acero"])
        c = ubicar_columnas(enc, MAPA_RECHAZOS, obligatorias=["orden", "acero", "fecha"])
        rechazos = []
        for i, r in enumerate(filas_r[h + 1:], start=h + 2):
            orden = texto(_celda(r, c["orden"]), 40)
            if orden is None:
                continue
            g = lambda campo, largo=None: texto(_celda(r, c.get(campo)), largo)
            rechazos.append((i, orden, g("zfer", 40), g("zhal", 40), g("acero", 80), g("vehiculo", 200),
                             g("plano", 40), fecha(_celda(r, c["fecha"]))))
    finally:
        wb.close()
    return {
        "cesar.aceros_consolidado": (COLUMNAS_CONSOLIDADO, consolidado),
        "cesar.aceros_rechazos": (COLUMNAS_RECHAZOS, rechazos),
    }, {}


COLUMNAS_PEDIDOS = ["fila", "servicio", "fecha", "material", "comentario", "marcos", "orden_trabajo", "valor"]


def leer_pedidos(ruta: Path):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = _filas(_hoja(wb, "HISTORIAL CORTES"))
    finally:
        wb.close()
    h, enc = buscar_encabezado(filas, ["Fecha", "Comentario"])
    c = ubicar_columnas(enc, {"servicio": ["Servicio"], "fecha": ["Fecha"], "material": ["Material"],
                              "comentario": ["Comentario"], "marcos": ["#Marcos", "Marcos", "# Marcos"],
                              "orden_trabajo": ["Orden trabajo", "Orden de trabajo"], "valor": ["Valor"]},
                        obligatorias=["fecha", "comentario", "marcos"])
    salida = []
    for i, r in enumerate(filas[h + 1:], start=h + 2):
        f = fecha(_celda(r, c["fecha"]))
        if f is None and _celda(r, c["comentario"]) is None:
            continue
        valor = numero(_celda(r, c.get("valor")))
        salida.append((i, texto(_celda(r, c.get("servicio")), 40), f, texto(_celda(r, c.get("material")), 80),
                       texto(_celda(r, c["comentario"]), 400), numero(_celda(r, c["marcos"])),
                       texto(_celda(r, c.get("orden_trabajo")), 400),
                       round(valor, 2) if valor is not None and abs(valor) < 1e15 else None))
    return {"cesar.pedidos_corte": (COLUMNAS_PEDIDOS, salida)}, {}


def leer_x5(ruta: Path):
    """Tabla FECHA | PIEZA | INVENTARIO MATRICERIA | INVENTARIO CALIDAD (la FECHA del Excel es =HOY(), no sirve)."""
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            filas = _filas(ws)
            for h, r in enumerate(filas[:20]):
                t = [norm(x) for x in r]
                if "PIEZA" in t and any("MATRICERIA" in x for x in t):
                    cp = t.index("PIEZA")
                    cm = next(i for i, x in enumerate(t) if "MATRICERIA" in x)
                    cc = next((i for i, x in enumerate(t) if "CALIDAD" in x), None)
                    salida = []
                    for i, fila in enumerate(filas[h + 1:], start=h + 2):
                        pieza = texto(_celda(fila, cp), 120)
                        if pieza is None or norm(pieza) == "TOTAL":
                            break
                        salida.append((i, pieza, numero(_celda(fila, cm)) or 0.0, numero(_celda(fila, cc)) or 0.0))
                    if salida:
                        return {"cesar.x5_inventario": (["fila", "pieza", "matriceria", "calidad"], salida)}, {"hoja": ws.title}
    finally:
        wb.close()
    raise ValueError("No encontré la tabla PIEZA / INVENTARIO MATRICERIA en ninguna hoja")


COLUMNAS_MOLDES = ["fila", "fecha_solicitud", "herramental", "codigo", "zhal", "pieza", "repeticion", "completo",
                   "vehiculo", "fecha_entrega", "fabricante", "estatus", "patron", "comentario", "tamano", "tiempo_h"]


def leer_fila_moldes(ruta: Path):
    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = None
        for ws in wb.worksheets:
            filas = _filas(ws)
            try:
                h, enc = buscar_encabezado(filas, ["Herramental", "Estatus", "Tiempo"])
                hoja = ws
                break
            except ValueError:
                continue
        if hoja is None:
            raise ValueError("No encontré la tabla de herramentales (Herramental / Estatus / Tiempo)")
        matriceros = None
        if "Cola" in wb.sheetnames:
            for r in _filas(wb["Cola"])[:10]:
                if r and norm(r[0]).startswith("MATRICEROS"):
                    matriceros = numero(r[1])
    finally:
        wb.close()
    c = ubicar_columnas(enc, {
        "fecha_solicitud": ["Fecha solicitud"], "herramental": ["Herramental"], "codigo": ["Codigo"], "zhal": ["Zhal"],
        "pieza": ["PIEZA"], "repeticion": ["Repeticion"], "completo": ["completo"], "vehiculo": ["Columna1", "Vehiculo"],
        "fecha_entrega": ["Fecha entrega"], "fabricante": ["Fabricante"], "estatus": ["Estatus"], "patron": ["Patron"],
        "comentario": ["Comentario"], "tamano": ["Tamano", "Tamaño"], "tiempo_h": ["Tiempo"],
    }, obligatorias=["herramental", "estatus", "tiempo_h", "fecha_entrega"])
    salida = []
    for i, r in enumerate(filas[h + 1:], start=h + 2):
        g = lambda campo, largo=None: texto(_celda(r, c.get(campo)), largo)
        if g("herramental") is None and g("estatus") is None:
            continue
        salida.append((i, fecha(_celda(r, c.get("fecha_solicitud"))), g("herramental", 40), g("codigo", 60), g("zhal", 60),
                       g("pieza", 40), g("repeticion", 20), g("completo", 80), g("vehiculo", 200),
                       fecha(_celda(r, c["fecha_entrega"])), g("fabricante", 120), g("estatus", 40), g("patron", 60),
                       g("comentario", 400), g("tamano", 20), numero(_celda(r, c["tiempo_h"]))))
    return {"cesar.fila_moldes": (COLUMNAS_MOLDES, salida)}, {"matriceros": matriceros, "hoja": hoja.title}


LECTORES = {
    "consolidado": leer_consolidado,
    "pedidos": leer_pedidos,
    "x5": leer_x5,
    "fila_moldes": leer_fila_moldes,
}
